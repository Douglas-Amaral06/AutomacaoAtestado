import hashlib
import io
import json
import os
import re
import secrets
import threading
import time
from contextlib import asynccontextmanager
from datetime import datetime
from datetime import timedelta
from pathlib import Path

from fastapi import BackgroundTasks, FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.middleware.trustedhost import TrustedHostMiddleware
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, RedirectResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook

from .database import BASE_DIR, UPLOAD_DIR, connect, initialize_database, new_operator_public_id
from .databricks_delivery import configured_delivery_service, prepare_processed_delivery, validate_prepared_delivery
from .maintenance import BACKUP_DIR, apply_retention, create_backup, detect_orphan_files, prune_backups
from .processing import QueueItemBusyError, add_log, process_queue_item, resume_pending_once, understandable_error
from .rate_limit import check_daily_quota, check_rate_limit
from .safe_errors import format_safe_error
from .export import safe_excel_value
from .bootstrap import bootstrap_users
from .config import positive_env_int
from .uploads import ingest_document, MAX_MULTIPART_BYTES, UploadBodyLimit
from .validation import document_type, normalize_cid, normalize_cpf, validation_summary
from .security import (SESSION_COOKIE, attempt_retry_after, create_session, current_user,
                       encrypt_totp, hash_password, is_login_blocked, login_keys,
                       permissions_for, record_login, is_attempt_blocked, require_csrf, require_permission,
                       trusted_client_ip, utc_now, verify_login_password)

templates = Jinja2Templates(directory=BASE_DIR / "app" / "templates")
_worker_stop = threading.Event()


class DeliveryCollisionError(RuntimeError):
    """Dois recebimentos locais apontam para o mesmo identificador contratual."""


def format_template_date(value, include_time: bool = False) -> str:
    """Formata datas do banco para exibição sem depender de JavaScript."""
    if value in (None, ""):
        return "—"

    parsed = value
    if isinstance(value, str):
        try:
            parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
        except ValueError:
            return value

    if not hasattr(parsed, "strftime"):
        return str(value)

    pattern = "%d/%m/%Y %H:%M" if include_time else "%d/%m/%Y"
    return parsed.strftime(pattern)


templates.env.filters["date_br"] = format_template_date
templates.env.filters["datetime_br"] = lambda value: format_template_date(value, include_time=True)


def queue_worker():
    while not _worker_stop.is_set():
        try: resume_pending_once()
        except Exception: pass
        _worker_stop.wait(20)


def maintenance_worker():
    while not _worker_stop.is_set():
        try:
            interval = max(1, int(os.getenv("BACKUP_INTERVAL_HOURS", "24"))) * 3600
            newest = max((item.stat().st_mtime for item in BACKUP_DIR.glob("*.zip")), default=0)
            if time.time() - newest >= interval:
                backup = create_backup()
                retention = apply_retention()
                removed = prune_backups()
                orphans = detect_orphan_files()
                add_log("info", "manutencao_concluida", "Backup automatico verificado e manutencao executada", {
                    "backup": backup.name, "retencao": retention, "backups_removidos": removed,
                    "arquivos_orfaos": len(orphans),
                })
        except Exception as error:
            try:
                safe_message, safe_details = format_safe_error(error)
                add_log("erro", "manutencao_falhou", safe_message, safe_details)
            except Exception: pass
        _worker_stop.wait(3600)


@asynccontextmanager
async def lifespan(_app):
    if os.getenv("APP_ENV", "development").lower() in {"pilot", "production"}:
        from .security import _app_secret
        _app_secret()
        if os.getenv("COOKIE_SECURE", "false").lower() != "true":
            raise RuntimeError("O piloto exige COOKIE_SECURE=true")
        if os.getenv("TRUST_CLOUDFLARE", "false").lower() != "false":
            raise RuntimeError("No piloto Render, TRUST_CLOUDFLARE deve ser false")
        if os.getenv("DATABRICKS_AUTH_MODE", "m2m").lower() != "m2m":
            raise RuntimeError("O piloto exige autenticação M2M")
    initialize_database()
    bootstrap_users()
    if os.getenv("APP_ENV", "development").lower() in {"pilot", "production"}:
        with connect() as connection:
            if not connection.execute("SELECT 1 FROM usuarios WHERE ativo=1 AND perfil='admin'").fetchone():
                raise RuntimeError("Configure um administrador ativo em BOOTSTRAP_USERS_JSON antes de iniciar o piloto")
    _worker_stop.clear()
    thread = threading.Thread(target=queue_worker, daemon=True, name="fila-atestados")
    maintenance = threading.Thread(target=maintenance_worker, daemon=True, name="manutencao-atestados")
    thread.start()
    if os.getenv("LOCAL_BACKUP_ENABLED", "false").lower() == "true":
        maintenance.start()
    try:
        yield
    finally:
        _worker_stop.set()


app = FastAPI(title="Recebimento Seguro de Atestados", docs_url=None, redoc_url=None, openapi_url=None, lifespan=lifespan)
app.add_middleware(TrustedHostMiddleware, allowed_hosts=[host.strip() for host in os.getenv("ALLOWED_HOSTS", "127.0.0.1,localhost").split(",") if host.strip()])
app.add_middleware(UploadBodyLimit)
app.mount("/static", StaticFiles(directory=BASE_DIR / "app" / "static", html=False), name="static")


@app.exception_handler(Exception)
async def unexpected_exception_handler(_request: Request, error: Exception):
    """Impede que exceções não tratadas exponham mensagens no FastAPI/Uvicorn."""
    safe_message, safe_details = format_safe_error(error)
    try:
        add_log("erro", "falha_interna", safe_message, safe_details)
    except Exception:
        pass
    return JSONResponse(
        status_code=500,
        content={"detail": {"codigo": "internal_error", "mensagem": safe_message}},
    )


def require_admin(user) -> None:
    if user["perfil"] != "admin": raise HTTPException(403, "Acesso exclusivo do administrador")


def parse_optional_boolean(value: str, field_name: str) -> bool | None:
    if value == "":
        return None
    if value == "true":
        return True
    if value == "false":
        return False
    raise HTTPException(422, f"Valor inválido para {field_name}.")


def review_context(request: Request, item, user, values: dict | None = None, validation: dict | None = None):
    record = dict(item)
    if values:
        record.update(values)
    if record.get("arquivo_hash"):
        with connect() as connection:
            record["possivel_repeticao"] = connection.execute(
                "SELECT COUNT(*) FROM atestados WHERE arquivo_hash=? AND id<>?",
                (record["arquivo_hash"], record.get("id") or 0),
            ).fetchone()[0] > 0
    validation = validation or validation_summary(record)
    return {
        "item": record, "user": user, "csrf": user["csrf_token"],
        "validation": validation,
    }


@app.middleware("http")
async def security_headers(request: Request, call_next):
    response = await call_next(request)
    response.headers.update({
        "X-Content-Type-Options": "nosniff", "X-Frame-Options": "SAMEORIGIN",
        "Referrer-Policy": "no-referrer", "Permissions-Policy": "camera=(), microphone=(), geolocation=()",
        "Content-Security-Policy": "default-src 'self'; script-src 'self'; style-src 'self' 'unsafe-inline'; img-src 'self' data:; connect-src 'self'; frame-src 'self'; frame-ancestors 'self'; form-action 'self'; base-uri 'none'; object-src 'none'",
        "Cache-Control": "no-store",
    })
    if os.getenv("COOKIE_SECURE", "false").lower() == "true":
        response.headers["Strict-Transport-Security"] = "max-age=31536000; includeSubDomains"
    return response


@app.middleware("http")
async def upload_limits(request: Request, call_next):
    """Autentica e limita uploads antes de o FastAPI analisar o multipart."""
    if request.method == "POST" and request.url.path == "/atestados/upload":
        try:
            user = current_user(request)
            require_permission(user, "upload")
            declared_size = request.headers.get("content-length", "")
            if not declared_size.isdigit():
                raise HTTPException(411, "Content-Length é obrigatório para uploads.")
            incoming_bytes = int(declared_size)
            if incoming_bytes <= 0:
                raise HTTPException(400, "Upload vazio.")
            if incoming_bytes > MAX_MULTIPART_BYTES:
                raise HTTPException(413, "Requisição de upload acima do limite permitido.")
        except HTTPException as error:
            return JSONResponse(
                status_code=error.status_code,
                content={"detail": error.detail},
                headers=error.headers,
            )
    return await call_next(request)


@app.middleware("http")
async def trusted_proxy_guard(request: Request, call_next):
    if os.getenv("TRUST_CLOUDFLARE", "false").casefold() == "true":
        try:
            trusted_client_ip(request)
        except HTTPException as error:
            return JSONResponse(
                status_code=error.status_code,
                content={"detail": error.detail},
            )
    return await call_next(request)


def web_user(request: Request):
    try: return current_user(request)
    except HTTPException: return None


def secure_request(request: Request) -> bool:
    if os.getenv("COOKIE_SECURE", "false").lower() == "true" or request.url.scheme == "https":
        return True
    return (
        os.getenv("TRUST_CLOUDFLARE", "false").lower() == "true"
        and request.headers.get("cf-visitor", "").find('"scheme":"https"') >= 0
    )


@app.get("/login", response_class=HTMLResponse)
def login_page(request: Request, erro: str = ""):
    if web_user(request): return RedirectResponse("/", 303)
    return templates.TemplateResponse(request=request, name="login.html", context={"erro": erro})


@app.post("/login")
def login(request: Request, usuario: str = Form(..., min_length=1, max_length=100), senha: str = Form(..., min_length=1, max_length=256)):
    ip_key, account_key = login_keys(request, usuario)
    if is_attempt_blocked(ip_key, 10, 15) or is_login_blocked(account_key):
        retry = max(attempt_retry_after(ip_key, 15), attempt_retry_after(account_key, 15))
        add_log("aviso", "login_bloqueado", "Tentativas excessivas de login bloqueadas", {"aguarde_segundos": retry})
        return RedirectResponse("/login?erro=Bloqueado+temporariamente.+Tente+novamente+mais+tarde", 303)
    with connect() as connection:
        user = connection.execute("SELECT * FROM usuarios WHERE usuario=? AND ativo=1", (usuario.strip(),)).fetchone()
    valid = verify_login_password(user, senha)
    record_login(ip_key, valid)
    record_login(account_key, valid)
    if not valid: return RedirectResponse("/login?erro=Credenciais+invalidas", 303)
    token, _, expires = create_session(user["id"], request)
    with connect() as connection: connection.execute("UPDATE usuarios SET ultimo_login=? WHERE id=?", (utc_now().isoformat(), user["id"]))
    response = RedirectResponse("/", 303)
    response.set_cookie(SESSION_COOKIE, token, httponly=True, secure=secure_request(request), samesite="strict", max_age=8*3600, path="/")
    add_log("info", "login", f"Login realizado por usuario #{user['id']}")
    return response


@app.post("/logout")
def logout(request: Request, csrf_token: str = Form(...)):
    user = current_user(request); require_csrf(request, user, csrf_token)
    with connect() as connection: connection.execute("DELETE FROM sessoes WHERE id=?", (user["session_id"],))
    response = RedirectResponse("/login", 303); response.delete_cookie(SESSION_COOKIE); return response


@app.get("/usuarios", response_class=HTMLResponse)
def users_page(request: Request):
    user=web_user(request)
    if not user:return RedirectResponse("/login",303)
    require_admin(user)
    with connect() as connection: rows=connection.execute("SELECT id,usuario,nome,perfil,ativo,criado_em,ultimo_login,operador_public_id FROM usuarios ORDER BY nome").fetchall()
    return templates.TemplateResponse(request=request,name="users.html",context={"usuarios":rows,"user":user,"csrf":user["csrf_token"],"provisioning_uri":None})


@app.post("/usuarios")
def create_user(request:Request,csrf_token:str=Form(..., max_length=100),usuario:str=Form(..., min_length=1, max_length=100),nome:str=Form(..., min_length=1, max_length=200),senha:str=Form(..., min_length=12, max_length=256),perfil:str=Form(..., pattern="^(admin|analista)$")):
    admin=current_user(request); require_csrf(request,admin,csrf_token); require_admin(admin)
    if perfil not in {"admin","analista"}:raise HTTPException(400,"Perfil invalido")
    secret=secrets.token_urlsafe(32)
    try:
        with connect() as connection: uid=connection.execute("INSERT INTO usuarios(usuario,nome,senha_hash,totp_secret_encrypted,perfil,operador_public_id) VALUES(?,?,?,?,?,?)",(usuario.strip(),nome.strip(),hash_password(senha),encrypt_totp(secret),perfil,new_operator_public_id())).lastrowid
    except Exception as error:
        raise HTTPException(400,"Usuario existente ou dados invalidos") from error
    add_log("info","usuario_criado",f"Usuario #{uid} criado pelo administrador #{admin['id']}")
    with connect() as connection: rows=connection.execute("SELECT id,usuario,nome,perfil,ativo,criado_em,ultimo_login,operador_public_id FROM usuarios ORDER BY nome").fetchall()
    return templates.TemplateResponse(request=request,name="users.html",context={"usuarios":rows,"user":admin,"csrf":admin["csrf_token"],"provisioning_uri":None})


@app.post("/usuarios/{user_id}/alternar")
def toggle_user(user_id:int,request:Request,csrf_token:str=Form(...)):
    admin=current_user(request);require_csrf(request,admin,csrf_token);require_admin(admin)
    if user_id==admin["id"]:raise HTTPException(400,"Nao desative a propria conta")
    with connect() as connection:
        target=connection.execute("SELECT perfil,ativo FROM usuarios WHERE id=?",(user_id,)).fetchone()
        if not target:raise HTTPException(404,"Usuario inexistente")
        if target["perfil"]=="admin" and target["ativo"]:
            active_admins=connection.execute("SELECT COUNT(*) FROM usuarios WHERE perfil='admin' AND ativo=1").fetchone()[0]
            if active_admins<=1:raise HTTPException(400,"O sistema deve manter ao menos um administrador ativo")
        connection.execute("UPDATE usuarios SET ativo=CASE ativo WHEN 1 THEN 0 ELSE 1 END WHERE id=?",(user_id,))
        connection.execute("DELETE FROM sessoes WHERE usuario_id=?",(user_id,))
    add_log("info","usuario_alternado",f"Status do usuario #{user_id} alterado por #{admin['id']}");return RedirectResponse("/usuarios",303)


@app.get("/", response_class=HTMLResponse)
def dashboard(request: Request, status: str = "", q: str = "", page: int = 1, per_page: int = 50):
    user = web_user(request)
    if not user: return RedirectResponse("/login", 303)
    page = max(1, page)
    per_page = min(100, max(10, per_page))
    offset = (page - 1) * per_page
    search = re.sub(r"[^A-Za-z0-9_-]", "", q.strip())[:120]
    conditions = []
    parameters = []
    if status:
        conditions.append("a.status=?")
        parameters.append(status)
    if search:
        conditions.append("instr(COALESCE(a.id_documento,''), ?) > 0")
        parameters.append(search)
    where = " WHERE " + " AND ".join(conditions) if conditions else ""
    with connect() as connection:
        total = connection.execute(
            "SELECT COUNT(*) FROM atestados a" + where, parameters
        ).fetchone()[0]
        rows = connection.execute(
            """SELECT a.*,u.nome revisor_nome,op.nome operador_envio_nome
               FROM atestados a
               LEFT JOIN usuarios u ON u.id=a.revisado_por
               LEFT JOIN usuarios op ON op.id=a.operador_envio_id"""
            + where + " ORDER BY a.id DESC LIMIT ? OFFSET ?",
            (*parameters, per_page, offset),
        ).fetchall()
        queue = connection.execute("SELECT status,COUNT(*) quantidade FROM fila_processamento GROUP BY status").fetchall()
        failed_items = connection.execute(
            """SELECT id,arquivo_original,status,tentativas,erro_amigavel,atualizado_em
               FROM fila_processamento WHERE status IN ('falhou','pausado_quota','aguardando_retentativa') ORDER BY atualizado_em DESC LIMIT 100"""
        ).fetchall()
        duplicate_hashes = {row[0] for row in connection.execute(
            "SELECT arquivo_hash FROM atestados WHERE arquivo_hash IS NOT NULL GROUP BY arquivo_hash HAVING COUNT(*)>1"
        ).fetchall()}
    atestados = []
    for row in rows:
        record = dict(row)
        record["possivel_repeticao"] = record.get("arquivo_hash") in duplicate_hashes
        atestados.append({**record, "validation": validation_summary(record)})
    failures = [
        {**dict(item), "mensagem": item["erro_amigavel"] or "A extração falhou. Reprocesse o documento ou consulte o suporte."}
        for item in failed_items
    ]
    return templates.TemplateResponse(request=request, name="dashboard.html", context={
        "atestados": atestados, "queue": queue, "failed_items": failures,
        "orphan_files": detect_orphan_files(), "permissions": permissions_for(user),
        "page": page, "per_page": per_page,
        "total_pages": max(1, (total + per_page - 1) // per_page), "status_filter": status,
        "search_query": search,
        "user": user, "csrf": user["csrf_token"],
    })


@app.get("/healthz")
def healthz():
    return {"status": "ok"}


@app.get("/atestados/novo", response_class=HTMLResponse)
def new_document(request: Request):
    user = web_user(request)
    if not user:
        return RedirectResponse("/login", 303)
    require_permission(user, "upload")
    return upload_page(request, user)


def upload_page(request, user, error=None, status_code=200):
    return templates.TemplateResponse(request=request, name="upload.html", context={
        "user": user, "csrf": user["csrf_token"], "error": error,
        "unit": os.getenv("DELIVERY_UNIT", "AUREA"), "pole": os.getenv("DELIVERY_POLO", "SP"),
        "reading_limit": positive_env_int("GEMINI_MAX_DOCUMENT_MB", 8),
    }, status_code=status_code)


def process_uploaded_item(queue_id):
    try:
        process_queue_item(queue_id)
    except Exception:
        pass  # A fila registra falhas e mantém o arquivo para retentativa.


@app.post("/atestados/upload")
def upload_document(request: Request, background_tasks: BackgroundTasks,
                    file: UploadFile = File(...), csrf_token: str = Form(..., max_length=100)):
    user = current_user(request)
    require_csrf(request, user, csrf_token)
    require_permission(user, "upload")
    check_rate_limit(str(user["id"]), positive_env_int("UPLOAD_RATE_LIMIT_PER_HOUR", 30), 3600)
    check_daily_quota(str(user["id"]), int(request.headers["content-length"]),
                      positive_env_int("UPLOAD_DAILY_QUOTA_MB", 300) * 1024 * 1024)
    try:
        queue_id = ingest_document(file, user["id"])
    except HTTPException as error:
        return upload_page(request, user, error.detail, error.status_code)
    background_tasks.add_task(process_uploaded_item, queue_id)
    return RedirectResponse(f"/envios/{queue_id}", 303)


@app.get("/envios/{queue_id}", response_class=HTMLResponse)
def upload_status(queue_id: int, request: Request):
    user = current_user(request)
    require_permission(user, "upload")
    with connect() as connection:
        item = connection.execute("SELECT * FROM fila_processamento WHERE id=?", (queue_id,)).fetchone()
        if not item or (item["operador_id"] != user["id"] and user["perfil"] != "admin"):
            raise HTTPException(404, "Envio inexistente")
        duplicate = connection.execute(
            "SELECT 1 FROM fila_processamento WHERE arquivo_hash=? AND id<>? LIMIT 1",
            (item["arquivo_hash"], queue_id),
        ).fetchone() is not None
    return templates.TemplateResponse(request=request, name="upload_status.html", context={
        "user": user, "csrf": user["csrf_token"], "item": item, "duplicate": duplicate,
    })


@app.get("/atestados/{record_id}", response_class=HTMLResponse)
def review_page(record_id: int, request: Request):
    user = web_user(request)
    if not user: return RedirectResponse("/login",303)
    require_permission(user, "review")
    with connect() as connection:
        row=connection.execute(
            """SELECT a.*,u.nome operador_envio_nome FROM atestados a
               LEFT JOIN usuarios u ON u.id=a.operador_envio_id WHERE a.id=?""",
            (record_id,),
        ).fetchone()
    if not row: raise HTTPException(404,"Registro inexistente")
    return templates.TemplateResponse(request=request,name="review.html",context=review_context(request, row, user))


@app.post("/atestados/{record_id}/revisar")
def review_document(
    record_id: int,
    request: Request,
    acao: str = Form(..., pattern="^(aprovar|rejeitar)$"),
    csrf_token: str = Form(..., max_length=100),
    versao_registro: str = Form(..., max_length=40),
    nome: str = Form("", max_length=200),
    cpf: str = Form("", max_length=20),
    cid: str = Form("", max_length=20),
    dias_afastamento: str = Form("", max_length=10),
    data_atestado: str = Form("", max_length=10),
    observacoes: str = Form("", max_length=4000),
    motivo_rejeicao: str = Form("", max_length=1000),
    matricula: str | None = Form(None, max_length=80),
    telefone: str | None = Form(None, max_length=80),
    email: str | None = Form(None, max_length=254),
    empresa: str | None = Form(None, max_length=200),
    tipo_documento: str = Form("", max_length=40),
    crm: str = Form("", max_length=30),
    crm_uf: str = Form("", max_length=2),
    assinado: str = Form("", max_length=5),
    carimbado: str = Form("", max_length=5),
):
    user=current_user(request); require_csrf(request,user,csrf_token)
    require_permission(user, "review")
    if acao not in {"aprovar","rejeitar"}: raise HTTPException(400,"Acao invalida")
    if acao=="rejeitar" and not motivo_rejeicao.strip(): raise HTTPException(400,"Informe o motivo")
    status="confirmado" if acao=="aprovar" else "rejeitado"
    with connect() as connection:
        before=connection.execute("SELECT * FROM atestados WHERE id=?",(record_id,)).fetchone()
    if not before: raise HTTPException(404,"Registro inexistente")
    employee = {key: value if value is not None else before[key] for key, value in {
        "matricula": matricula, "telefone": telefone, "email": email, "empresa": empresa,
    }.items()}
    enrichment_status = before["status_enriquecimento"]
    signed_value = parse_optional_boolean(assinado, "assinatura")
    stamped_value = parse_optional_boolean(carimbado, "carimbo")
    reviewed = {"nome":nome.strip() or None,"cpf":normalize_cpf(cpf),"cid":normalize_cid(cid),"dias_afastamento":dias_afastamento.strip() or None,"data_atestado":data_atestado.strip() or None,"tipo_documento":document_type(tipo_documento) or tipo_documento.strip() or None,"status_enriquecimento":enrichment_status,"crm":crm.strip() or None,"crm_uf":crm_uf.strip().upper() or None,"assinado":signed_value,"carimbado":stamped_value}
    validation = validation_summary(reviewed)
    if acao == "aprovar" and validation["errors"]:
        add_log("aviso", "aprovacao_bloqueada_validacao", f"Atestado #{record_id} exige correcao", {"erros": validation["errors"]})
        return templates.TemplateResponse(request=request, name="review.html", context=review_context(request, before, user, {**reviewed, **employee, "observacoes": observacoes, "motivo_rejeicao": motivo_rejeicao}, validation), status_code=422)
    days = int(dias_afastamento) if dias_afastamento.strip().isdigit() else None
    reviewed["dias_afastamento"] = days
    reservation = utc_now().isoformat()
    lease_cutoff = (utc_now() - timedelta(minutes=30)).isoformat()
    # Salva a edição antes da rede. A reserva é uma comparação atômica da versão
    # e do estado, inclusive quando outra pessoa reabre a tela durante a entrega.
    fields = {**reviewed, **employee, "observacoes": observacoes.strip() or None,
              "motivo_rejeicao": motivo_rejeicao.strip() or None,
              "revisado_por": user["id"], "revisado_em": reservation,
              "status": "pendente" if acao == "aprovar" else "rejeitado",
              "status_entrega": "entregando" if acao == "aprovar" else "rejeitado"}
    with connect() as connection:
        assignments = ",".join(f"{key}=?" for key in fields)
        cursor = connection.execute(
            f"""UPDATE atestados SET {assignments} WHERE id=?
                AND COALESCE(revisado_em,criado_em)=?
                AND status<>'confirmado' AND COALESCE(status_entrega,'')<>'entregue_volume'
                AND (COALESCE(status_entrega,'')<>'entregando' OR revisado_em<?)""",
            (*fields.values(), record_id, versao_registro, lease_cutoff),
        )
    if cursor.rowcount != 1:
        raise HTTPException(409, "Este atestado foi alterado, já foi confirmado ou está sendo enviado. Reabra a tela.")
    if acao == "aprovar":
        try:
            mode = os.getenv("DELIVERY_MODE", "disabled").strip().lower()
            if mode != "databricks" and os.getenv("APP_ENV", "development").lower() not in {"development", "local", "test"}:
                raise RuntimeError("O piloto exige entrega real ao Databricks")
            delivery_service = configured_delivery_service()
            if delivery_service is None:
                raise RuntimeError("A entrega não está habilitada")
            with connect() as connection:
                queue_item = connection.execute(
                    """SELECT q.*,u.operador_public_id FROM fila_processamento q
                       LEFT JOIN usuarios u ON u.id=q.operador_id
                       WHERE q.atestado_id=? ORDER BY q.id DESC LIMIT 1""", (record_id,),
                ).fetchone()
            document_path = UPLOAD_DIR / before["arquivo_salvo"]
            if not queue_item or not document_path.is_file():
                raise RuntimeError("Original ou metadados indisponíveis")
            content = document_path.read_bytes()
            if hashlib.sha256(content).hexdigest() != queue_item["arquivo_hash"]:
                raise RuntimeError("Integridade do original não confere")
            original = json.loads(before["dados_originais"] or "{}")
            approved = {**original, **reviewed, "observacoes": observacoes.strip() or None,
                        "is_atestado": True, "revisao_humana": {
                            "status": "aprovado", "operador_id": user["operador_public_id"],
                            "data_revisao": reservation,
                        }}
            prepared = prepare_processed_delivery(queue_item, approved, content)
            validate_prepared_delivery(prepared)
            with connect() as connection:
                connection.execute("BEGIN IMMEDIATE")
                collision = connection.execute(
                    "SELECT 1 FROM atestados WHERE id_documento=? AND id<>?",
                    (prepared.payload["id_documento"], record_id),
                ).fetchone()
                if collision:
                    raise DeliveryCollisionError()
                connection.execute("UPDATE atestados SET id_documento=? WHERE id=? AND revisado_em=?",
                                   (prepared.payload["id_documento"], record_id, reservation))
            delivery_service.deliver(prepared)
            real_delivery = mode == "databricks"
            status = "confirmado" if real_delivery else "pendente"
            with connect() as connection:
                cursor = connection.execute(
                    """UPDATE atestados SET status=?,status_entrega=?,id_documento=?
                       WHERE id=? AND revisado_em=? AND status_entrega='entregando'""",
                    (status, "entregue_volume" if real_delivery else "simulado_local",
                     prepared.payload["id_documento"], record_id, reservation),
                )
                if cursor.rowcount != 1:
                    raise RuntimeError("Reserva de entrega perdida")
        except Exception as error:
            safe_message, details = format_safe_error(error)
            add_log("erro", "entrega_falhou", safe_message, {**details, "atestado_id": record_id})
            with connect() as connection:
                connection.execute(
                    """UPDATE atestados SET status='pendente',status_entrega='falha_entrega'
                       WHERE id=? AND revisado_em=? AND status_entrega='entregando'""", (record_id, reservation),
                )
                saved = connection.execute("SELECT * FROM atestados WHERE id=?", (record_id,)).fetchone()
            context = review_context(request, saved, user)
            context["delivery_error"] = (
                "Não foi possível confirmar a entrega ao Databricks. Suas correções foram salvas. "
                "O registro continua pendente; tente aprovar novamente. Referência: " + details["correlation_id"]
            )
            if isinstance(error, DeliveryCollisionError):
                context["delivery_error"] = "Outro envio já utiliza este ID Databricks. Suas correções foram salvas. Envie novamente o documento para criar um novo recebimento."
            return templates.TemplateResponse(request=request, name="review.html", context=context, status_code=503)
    changed_values = {**locals(), "assinado": signed_value, "carimbado": stamped_value, "crm_uf": reviewed["crm_uf"]}
    add_log("info","revisao",f"Atestado #{record_id} {status} por usuario #{user['id']}",{"campos_alterados":[k for k in ("nome","cpf","cid","dias_afastamento","data_atestado","observacoes","crm","crm_uf","assinado","carimbado") if str(before[k] if before[k] is not None else "")!=str(changed_values[k] if changed_values[k] is not None else "")]})
    return RedirectResponse("/",303)


@app.post("/atestados/{record_id}/excluir")
def delete_document(record_id:int, request:Request, csrf_token:str=Form(...)):
    user=current_user(request); require_csrf(request,user,csrf_token)
    require_permission(user, "delete")
    with connect() as connection:
        connection.execute("BEGIN IMMEDIATE")
        item=connection.execute("SELECT arquivo_salvo,arquivo_hash,status_entrega FROM atestados WHERE id=?",(record_id,)).fetchone()
        if not item:raise HTTPException(404,"Registro inexistente")
        if item["status_entrega"] == "entregando":
            raise HTTPException(409, "Aguarde a entrega antes de excluir o registro.")
        connection.execute("DELETE FROM fila_processamento WHERE atestado_id=?",(record_id,))
        connection.execute("DELETE FROM atestados WHERE id=?",(record_id,))
        still_used=connection.execute("SELECT COUNT(*) FROM atestados WHERE arquivo_salvo=?",(item["arquivo_salvo"],)).fetchone()[0]
    if not still_used:
        (UPLOAD_DIR/item["arquivo_salvo"]).unlink(missing_ok=True)
    add_log("aviso","atestado_excluido",f"Atestado #{record_id} e arquivo removidos pelo usuario #{user['id']}")
    return RedirectResponse("/",303)


@app.get("/atestados/{record_id}/arquivo")
def download_original(record_id:int, request:Request):
    user = web_user(request)
    if not user: return RedirectResponse("/login",303)
    require_permission(user, "review")
    with connect() as connection: row=connection.execute("SELECT arquivo_original,arquivo_salvo FROM atestados WHERE id=?",(record_id,)).fetchone()
    if not row: raise HTTPException(404,"Registro inexistente")
    if not (UPLOAD_DIR / row["arquivo_salvo"]).is_file():
        raise HTTPException(404, "Original local indisponível. Documentos confirmados permanecem no Databricks.")
    return FileResponse(UPLOAD_DIR/row["arquivo_salvo"],filename=row["arquivo_original"],headers={"Content-Disposition":f"inline; filename=\"documento-{record_id}{Path(row['arquivo_original']).suffix}\""})


@app.get("/logs",response_class=HTMLResponse)
def logs_page(request:Request):
    user=web_user(request)
    if not user:return RedirectResponse("/login",303)
    require_admin(user)
    with connect() as connection: rows=connection.execute("SELECT * FROM logs ORDER BY id DESC LIMIT 500").fetchall()
    return templates.TemplateResponse(request=request,name="logs.html",context={"logs":rows,"user":user,"csrf":user["csrf_token"]})


@app.post("/fila/retomar")
def resume_queue(request:Request,csrf_token:str=Form(...)):
    user=current_user(request); require_csrf(request,user,csrf_token)
    require_admin(user)
    with connect() as connection: connection.execute("UPDATE fila_processamento SET status='aguardando_retentativa',disponivel_em=? WHERE status IN ('pausado_quota','falhou')",(utc_now().isoformat(),))
    add_log("info","fila_retomada",f"Fila retomada por usuario #{user['id']}"); return RedirectResponse("/",303)


@app.post("/fila/{queue_id}/reprocessar")
def reprocess_queue_item(queue_id: int, request: Request, csrf_token: str = Form(...)):
    user = current_user(request)
    require_csrf(request, user, csrf_token)
    require_permission(user, "reprocess")
    with connect() as connection:
        connection.execute("BEGIN IMMEDIATE")
        item = connection.execute(
            "SELECT status,atestado_id FROM fila_processamento WHERE id=?", (queue_id,)
        ).fetchone()
        if not item:
            raise HTTPException(404, "Item da fila inexistente")
        if item["atestado_id"] or item["status"] not in {"falhou", "pausado_quota", "aguardando_retentativa"}:
            raise HTTPException(409, "Somente extrações com falha podem ser reprocessadas")
        connection.execute(
            """UPDATE fila_processamento SET status='aguardando_retentativa',tentativas=0,
               disponivel_em=NULL,ultimo_erro=NULL,erro_amigavel=NULL,
               lock_token=NULL,lock_expires_em=NULL,atualizado_em=? WHERE id=?""",
            (utc_now().isoformat(), queue_id),
        )
    add_log("info", "reprocessamento_solicitado", f"Fila #{queue_id} reenviada por usuario #{user['id']}")
    try:
        process_queue_item(queue_id)
    except QueueItemBusyError as error:
        raise HTTPException(409, understandable_error(error)) from error
    except Exception:
        pass
    return RedirectResponse("/", 303)


@app.post("/fila/{queue_id}/excluir")
def delete_queue_item(
    queue_id: int,
    request: Request,
    csrf_token: str = Form(...),
):
    user = current_user(request)

    require_csrf(request, user, csrf_token)
    require_permission(user, "delete")

    with connect() as connection:
        connection.execute("BEGIN IMMEDIATE")
        item = connection.execute(
            """
            SELECT
                id,
                arquivo_hash,
                arquivo_original,
                arquivo_salvo,
                status,
                atestado_id
            FROM fila_processamento
            WHERE id = ?
            """,
            (queue_id,),
        ).fetchone()

        if not item:
            raise HTTPException(
                status_code=404,
                detail="Item da fila inexistente",
            )

        if item["atestado_id"]:
            raise HTTPException(
                status_code=409,
                detail=(
                    "Esta extração já possui um atestado associado "
                    "e não pode ser removida por esta ação."
                ),
            )

        allowed_statuses = {
            "falhou",
            "pausado_quota",
            "aguardando_retentativa",
        }

        if item["status"] not in allowed_statuses:
            raise HTTPException(
                status_code=409,
                detail="Somente extrações com falha podem ser excluídas.",
            )

        arquivo_salvo = item["arquivo_salvo"]
        arquivo_hash = item["arquivo_hash"]
        arquivo_original = item["arquivo_original"]

        connection.execute(
            """
            DELETE FROM fila_processamento
            WHERE id = ?
            """,
            (queue_id,),
        )

        queue_file_uses = connection.execute(
            """
            SELECT COUNT(*)
            FROM fila_processamento
            WHERE arquivo_salvo = ?
            """,
            (arquivo_salvo,),
        ).fetchone()[0]

        document_file_uses = connection.execute(
            """
            SELECT COUNT(*)
            FROM atestados
            WHERE arquivo_salvo = ?
            """,
            (arquivo_salvo,),
        ).fetchone()[0]

    # Só remove fisicamente o arquivo se nenhum registro ainda o utiliza.
    if (
        arquivo_salvo
        and queue_file_uses == 0
        and document_file_uses == 0
    ):
        (UPLOAD_DIR / arquivo_salvo).unlink(missing_ok=True)

    add_log(
        "aviso",
        "extracao_falha_excluida",
        f"Fila #{queue_id} removida pelo usuario #{user['id']}",
        {
            "arquivo": arquivo_original,
            "arquivo_hash": arquivo_hash,
        },
    )

    return RedirectResponse("/", status_code=303)


@app.get("/relatorios",response_class=HTMLResponse)
def reports(request:Request):
    user=web_user(request)
    if not user:return RedirectResponse("/login",303)
    require_permission(user, "reports")
    with connect() as connection:
        summary=connection.execute("SELECT COUNT(*) total,SUM(CASE WHEN status='pendente' THEN 1 ELSE 0 END) pendentes,SUM(CASE WHEN status='confirmado' THEN 1 ELSE 0 END) confirmados,SUM(CASE WHEN status='rejeitado' THEN 1 ELSE 0 END) rejeitados,COALESCE(SUM(CASE WHEN status='confirmado' THEN dias_afastamento ELSE 0 END),0) dias FROM atestados").fetchone()
        monthly=connection.execute("SELECT substr(COALESCE(data_atestado,criado_em),1,7) mes,COUNT(*) quantidade,COALESCE(SUM(dias_afastamento),0) dias FROM atestados GROUP BY mes ORDER BY mes DESC LIMIT 12").fetchall()
        timing=connection.execute("SELECT ROUND(AVG((julianday(revisado_em)-julianday(criado_em))*24),2) horas FROM atestados WHERE revisado_em IS NOT NULL").fetchone()
    return templates.TemplateResponse(request=request,name="reports.html",context={"summary":summary,"monthly":monthly,"timing":timing,"user":user,"csrf":user["csrf_token"]})


@app.get("/exportar.xlsx")
def export_xlsx(request:Request):
    user = web_user(request)
    if not user:return RedirectResponse("/login",303)
    require_permission(user, "export")
    headers = ["Matricula","Nome","CPF","Telefone","E-mail","Empresa","Tipo de documento","CRM","UF CRM","CID","Assinado","Carimbado","Dias","Data","Status","Status do enriquecimento","Observacoes","Recebido em","Revisado em"]
    query = "SELECT matricula,nome,cpf,telefone,email,empresa,tipo_documento,crm,crm_uf,cid,assinado,carimbado,dias_afastamento,data_atestado,status,status_enriquecimento,observacoes,criado_em,revisado_em FROM atestados ORDER BY id"
    workbook = Workbook(write_only=True)
    sheet_number = 1
    sheet = workbook.create_sheet("Atestados")
    sheet.append(headers)
    rows_in_sheet = 1
    with connect() as connection:
        cursor = connection.execute(query)
        while batch := cursor.fetchmany(1000):
            for row in batch:
                if rows_in_sheet >= 1_048_576:
                    sheet_number += 1
                    sheet = workbook.create_sheet(f"Atestados {sheet_number}")
                    sheet.append(headers)
                    rows_in_sheet = 1
                sheet.append([safe_excel_value(value) for value in row])
                rows_in_sheet += 1
    output = io.BytesIO()
    try:
        workbook.save(output)
    finally:
        workbook.close()
    output.seek(0)
    return StreamingResponse(output,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":"attachment; filename=atestados.xlsx"})
